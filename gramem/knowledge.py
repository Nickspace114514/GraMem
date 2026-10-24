"""Durable originals and source-located knowledge, sharing GraMem's scope and exclusions."""

import hashlib
import io
import json
import mimetypes
import os
import re
import shutil
import subprocess
import tempfile
import time
import zipfile
from pathlib import Path

from .config import Principal
from .documents import chunks
from .engine import SECRET, terms
from .store import Store, decode
from .tokens import compact, limit

SUFFIXES = {
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".pdf",
    ".docx",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".gif",
    ".tif",
    ".tiff",
    ".bmp",
}
IMAGES = SUFFIXES - {".txt", ".md", ".csv", ".json", ".pdf", ".docx"}
MAX_BYTES = 20_000_000
MAX_PAGES = 200
MAX_TEXT = 2_000_000


class _LeaseLost(Exception):
    """The processing lease moved to another worker before commit."""


class _Requeue(Exception):
    """Unrelated memory controls changed during parsing; parse again."""

    def __init__(self, epoch):
        super().__init__("memory controls changed during parsing")
        self.epoch = epoch


def safe_name(name):
    name = str(name).replace("\\", "/").rsplit("/", 1)[-1]
    name = "".join(c for c in name if ord(c) >= 32 and ord(c) != 127).strip()
    if not name or len(name) > 200 or Path(name).suffix.lower() not in SUFFIXES:
        raise ValueError("unsupported or invalid filename")
    return name


def _run(args, timeout=60):
    try:
        result = subprocess.run(
            args, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=timeout, check=True
        )
        return result.stdout.decode("utf-8")
    except (OSError, subprocess.SubprocessError, UnicodeError) as exc:
        raise RuntimeError("local parser failed or timed out") from exc


def _ocr(path, languages="chi_sim+eng"):
    if not shutil.which("tesseract"):
        raise ImportError("tesseract is not installed")
    langs = set(_run(["tesseract", "--list-langs"]).splitlines())
    selected = languages.split("+")
    if not selected or not set(selected).issubset(langs):
        raise ImportError("OCR language packs are not installed")
    return _run(["tesseract", str(path), "stdout", "-l", "+".join(selected)])


def extract(data, suffix, *, ocr_languages="chi_sim+eng"):
    """Return (page, method, text), physical page count, and parser limitations."""
    if suffix in {".txt", ".md", ".csv", ".json"}:
        return [(1, "text", data.decode("utf-8-sig"))], 1, []
    with tempfile.TemporaryDirectory(prefix="gramem-parse-") as folder:
        path = Path(folder) / ("input" + suffix)
        path.write_bytes(data)
        if suffix == ".pdf":
            from pypdf import PdfReader

            pdf = PdfReader(io.BytesIO(data))
            if pdf.is_encrypted:
                raise ValueError("encrypted PDF requires an unlocked copy")
            if len(pdf.pages) > MAX_PAGES:
                raise ValueError("document exceeds 200 pages")
            units, warnings = [], []
            ocr_pages = 0
            deadline = time.monotonic() + 600
            for i, page in enumerate(pdf.pages, 1):
                if time.monotonic() > deadline:
                    warnings.append(f"remaining pages from {i}: parser time budget exceeded")
                    break
                text = page.extract_text() or ""
                method = "text"
                if not text.strip():
                    if ocr_pages >= 30:
                        warnings.append(f"page {i}: OCR limit of 30 pages reached")
                    elif shutil.which("pdftoppm") and shutil.which("tesseract"):
                        ocr_pages += 1
                        image = Path(folder) / "page"
                        _run(
                            [
                                "pdftoppm",
                                "-f",
                                str(i),
                                "-l",
                                str(i),
                                "-scale-to",
                                "2400",
                                "-singlefile",
                                "-png",
                                str(path),
                                str(image),
                            ]
                        )
                        try:
                            text, method = _ocr(image.with_suffix(".png"), ocr_languages), "ocr"
                        except ImportError:
                            warnings.append(f"page {i}: OCR language packs unavailable")
                    else:
                        warnings.append(f"page {i}: OCR unavailable")
                units.append((i, method, text))
                if sum(len(u[2]) for u in units) > MAX_TEXT:
                    raise ValueError("extracted text exceeds 2 million characters")
            return units, len(pdf.pages), warnings
        if suffix == ".docx":
            with zipfile.ZipFile(io.BytesIO(data)) as package:
                if sum(f.file_size for f in package.infolist()) > 100_000_000:
                    raise ValueError("expanded document exceeds 100 MB")
            from docx import Document

            doc = Document(io.BytesIO(data))
            # Word pages depend on layout. A section number is not a page number.
            text = "\n".join(
                [p.text for p in doc.paragraphs]
                + [" | ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows]
            )
            return [(0, "text", text)], 0, []
        from PIL import Image, ImageOps

        with Image.open(io.BytesIO(data)) as image:
            if image.width * image.height > 25_000_000:
                raise ValueError("image exceeds 25 million pixels")
            image.seek(0)
            normalized = ImageOps.exif_transpose(image).convert("RGB")
            target = Path(folder) / "image.png"
            normalized.save(target)
            frames = getattr(image, "n_frames", 1)
        warnings = ["only first image frame indexed"] if frames > 1 else []
        try:
            return [(1, "ocr", _ocr(target, ocr_languages))], 1, warnings
        except ImportError:
            return [], 1, [*warnings, "OCR unavailable"]


class Knowledge:
    def __init__(self, engine):
        self.engine = engine
        self.store = engine.store
        self.root = self.store.path.parent / "knowledge"

    def blob(self, id):
        if not re.fullmatch(r"[a-f0-9]{64}", id):
            raise ValueError("invalid document id")
        return self.root / "blobs" / id[:2] / (id + ".blob")

    def _allowed(self, db, p, row):
        if p.scope == "group" or row["status"] in {"redacted", "blocked"}:
            return False
        if not self.engine._allowed(db, p, "document:" + row["id"]):
            return False
        return not db.execute(
            """SELECT 1 FROM knowledge_chunks k JOIN memories m ON m.id=k.memory_id
            JOIN exclusions x ON x.principal=m.principal AND x.slot IN (m.slot,'*')
            WHERE k.document_id=? AND x.state!='active' LIMIT 1""",
            (row["id"],),
        ).fetchone()

    def _get(self, db, p, id):
        row = db.execute(
            "SELECT * FROM knowledge_documents WHERE id=? AND principal=?", (id, p.key)
        ).fetchone()
        if row is None or not self._allowed(db, p, row):
            raise KeyError("document not available in this scope")
        return dict(row)

    def _info(self, db, row):
        result = dict(row)
        result.pop("principal", None)
        result.pop("epoch", None)
        result.pop("lease", None)
        result["origin"] = json.loads(result["origin"])
        result["chunks"] = db.execute(
            "SELECT count(*) FROM knowledge_chunks WHERE document_id=?", (row["id"],)
        ).fetchone()[0]
        result["original_available"] = self.blob(row["id"]).is_file()
        return result

    def ingest(
        self,
        p,
        path=None,
        *,
        data=None,
        filename=None,
        title="",
        description="",
        origin=None,
        defer=False,
    ):
        if p.scope == "group":
            raise PermissionError("knowledge documents are private")
        if path is not None:
            path = Path(path)
            if not path.is_file() or path.stat().st_size > MAX_BYTES:
                raise ValueError("file unavailable or exceeds 20 MB")
            filename, data = filename or path.name, path.read_bytes()
        name = safe_name(filename)
        if not isinstance(data, bytes) or not 0 < len(data) <= MAX_BYTES:
            raise ValueError("empty file or file exceeds 20 MB")
        if (
            not isinstance(title, str)
            or not isinstance(description, str)
            or len(title) > 200
            or len(description) > 2000
        ):
            raise ValueError("invalid title or description")
        if SECRET.search(title + description):
            raise ValueError("secret material is not allowed")
        suffix = Path(name).suffix.lower()
        if suffix in {".txt", ".md", ".csv", ".json"} and SECRET.search(data.decode("utf-8-sig")):
            raise ValueError("secret material is not allowed")
        digest = hashlib.sha256(data).hexdigest()
        id = hashlib.sha256((p.key + "\x1fasset\x1f" + digest).encode()).hexdigest()
        family = hashlib.sha256((p.key + "\x1fname\x1f" + name.casefold()).encode()).hexdigest()
        target = self.blob(id)
        with self.store.knowledge_dispatch():
            with self.store.transaction() as db:
                if db.execute(
                    "SELECT 1 FROM exclusions WHERE principal=? AND state!='active'", (p.key,)
                ).fetchone():
                    raise PermissionError("knowledge import paused while exclusions are active")
                old = db.execute("SELECT * FROM knowledge_documents WHERE id=?", (id,)).fetchone()
                if old:
                    if not self._allowed(db, p, old):
                        raise PermissionError("forgotten document cannot be restored")
                    return {
                        **self._info(db, old),
                        "source": id,
                        "duplicate": True,
                        "ids": [
                            r[0]
                            for r in db.execute(
                                "SELECT memory_id FROM knowledge_chunks WHERE document_id=?", (id,)
                            )
                        ],
                    }
                previous = db.execute(
                    "SELECT id,version FROM knowledge_documents WHERE principal=? AND family=? ORDER BY version DESC LIMIT 1",
                    (p.key, family),
                ).fetchone()
                version = previous["version"] + 1 if previous else 1
                target.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
                os.chmod(target.parent.parent, 0o700)
                os.chmod(target.parent, 0o700)
                os.chmod(self.root, 0o700)
                fd, tmp = tempfile.mkstemp(dir=target.parent, prefix=".upload-")
                try:
                    with os.fdopen(fd, "wb") as out:
                        out.write(data)
                        out.flush()
                        os.fsync(out.fileno())
                    os.replace(tmp, target)
                    directory_fd = os.open(target.parent, os.O_RDONLY)
                    try:
                        os.fsync(directory_fd)
                    finally:
                        os.close(directory_fd)
                    now = self.engine.clock()
                    db.execute(
                        """INSERT INTO knowledge_documents
                        (id,principal,name,title,description,sha256,size,suffix,mime,family,version,previous,status,created,updated,epoch,origin)
                        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                        (
                            id,
                            p.key,
                            name,
                            title or name,
                            description,
                            digest,
                            len(data),
                            suffix,
                            mimetypes.guess_type(name)[0] or "application/octet-stream",
                            family,
                            version,
                            previous["id"] if previous else None,
                            "queued",
                            now,
                            now,
                            self.engine._epoch(db, p),
                            compact(origin or {}),
                        ),
                    )
                    Store.changed(db)
                except BaseException:
                    target.unlink(missing_ok=True)
                    raise
                finally:
                    Path(tmp).unlink(missing_ok=True)
        if not defer:
            self.process_one(id=id)
        try:
            result = self.get(p, id)
        except KeyError as exc:
            raise PermissionError("document blocked by memory controls") from exc
        return {
            **result,
            "source": id,
            "duplicate": False,
            "ids": [c["memory_id"] for c in result["locations"]],
        }

    def process_one(self, *, id=None):
        now = self.engine.clock()
        with self.store.transaction() as db:
            row = db.execute(
                """SELECT * FROM knowledge_documents WHERE
                (status='queued' OR (status='processing' AND lease<?)) AND (? IS NULL OR id=?)
                ORDER BY created LIMIT 1""",
                (now, id, id),
            ).fetchone()
            if not row:
                return False
            row = dict(row)
            lease = now + 1800
            db.execute(
                "UPDATE knowledge_documents SET status='processing',lease=?,attempts=attempts+1 WHERE id=?",
                (lease, row["id"]),
            )
        parts = row["principal"].split("\x1f")
        p = Principal(*parts, parts[0] if parts[1] == "owner" else parts[2])
        try:
            with self.store.knowledge_dispatch():
                with self.store.transaction() as db:
                    self._get(db, p, row["id"])
                    data = self.blob(row["id"]).read_bytes()
            units, pages, warnings = extract(
                data, row["suffix"], ocr_languages=self.engine.config.knowledge_ocr_languages
            )
            if sum(len(u[2]) for u in units) > MAX_TEXT:
                raise ValueError("extracted text exceeds 2 million characters")
            text = "\n".join(t for _, _, t in units)
            if SECRET.search(text):
                self.engine.exclude(p, "document:" + row["id"])
                return True
            pieces = [
                (page, method, piece)
                for page, method, text in units
                for piece in chunks(text)
                if piece.strip()
            ]
            if row["description"]:
                pieces.insert(0, (0, "description", row["description"]))
            if len(pieces) > 5000:
                raise ValueError("document exceeds 5000 chunks")
            with self.store.transaction() as db:
                current = self._get(db, p, row["id"])
                if current["status"] != "processing" or current["lease"] != lease:
                    # Another worker reclaimed the lease; it owns the outcome.
                    raise _LeaseLost()
                epoch = self.engine._epoch(db, p)
                if current["epoch"] != epoch:
                    # A deletion may target a value inside this not-yet-indexed text,
                    # so keep the fail-closed block for deletions. _get() above already
                    # rejects documents whose own slot, '*', or chunk slots are
                    # excluded; suppressions and reactivations of other slots only
                    # require parsing again under the current policy.
                    if db.execute(
                        "SELECT 1 FROM exclusions WHERE principal=? AND state='deleted' AND updated>=? LIMIT 1",
                        (p.key, current["created"]),
                    ).fetchone():
                        raise PermissionError(
                            "memory deletion applied during parsing"
                        )
                    raise _Requeue(epoch)
                db.execute(
                    "INSERT OR IGNORE INTO sources(id,principal,session,kind,content,created,status,epoch) VALUES(?,?,'','document',?,?,'done',?)",
                    (row["id"], p.key, text, now, current["epoch"]),
                )
                if pieces:
                    older = db.execute(
                        """WITH RECURSIVE affected(id) AS (
                        SELECT k.memory_id FROM knowledge_chunks k
                        JOIN knowledge_documents d ON d.id=k.document_id
                        WHERE d.principal=? AND d.family=? AND d.version<?
                        UNION
                        SELECT m.id FROM memories m JOIN json_each(m.metadata,'$.targets') j
                        JOIN affected a ON a.id=j.value WHERE m.principal=?
                        ) SELECT m.id FROM affected a JOIN memories m ON m.id=a.id
                        WHERE m.principal=? AND m.status IN ('active','candidate')""",
                        (p.key, row["family"], row["version"], p.key, p.key),
                    ).fetchall()
                    for old in older:
                        self.engine._invalidate(db, p, old["id"])
                newer = db.execute(
                    "SELECT 1 FROM knowledge_documents WHERE principal=? AND family=? AND version>? AND status IN ('ready','partial')",
                    (p.key, row["family"], row["version"]),
                ).fetchone()
                for ordinal, (page, method, piece) in enumerate(pieces):
                    memory_id = self.engine._insert(
                        db,
                        p,
                        piece,
                        "knowledge",
                        f"document:{row['id']}:{ordinal}",
                        source=row["id"],
                        metadata={
                            "document_id": row["id"],
                            "document_name": row["name"],
                            "document_source": row["id"],
                            "chunk": ordinal,
                            "page": page,
                            "method": method,
                            "document_version": row["version"],
                        },
                    )
                    db.execute("UPDATE memories SET decay=0 WHERE id=?", (memory_id,))
                    if newer:
                        self.engine._invalidate(db, p, memory_id)
                    db.execute(
                        "INSERT OR IGNORE INTO knowledge_chunks VALUES(?,?,?,?,?)",
                        (row["id"], memory_id, ordinal, page, method),
                    )
                status = (
                    "partial"
                    if warnings and pieces
                    else "ready"
                    if pieces
                    else "needs_parser"
                    if warnings
                    else "empty"
                )
                db.execute(
                    "UPDATE knowledge_documents SET status=?,pages=?,error=?,updated=?,lease=0 WHERE id=?",
                    (status, pages, "; ".join(warnings)[:1000], self.engine.clock(), row["id"]),
                )
                Store.changed(db)
        except _LeaseLost:
            return True
        except _Requeue as requeue:
            with self.store.transaction() as db:
                db.execute(
                    """UPDATE knowledge_documents SET status='queued',epoch=?,lease=0,
                    attempts=MAX(attempts-1,0),updated=? WHERE id=? AND status='processing' AND lease=?""",
                    (requeue.epoch, self.engine.clock(), row["id"], lease),
                )
            return True
        except Exception as exc:
            status = (
                "blocked"
                if isinstance(exc, PermissionError)
                else "needs_parser"
                if isinstance(exc, ImportError)
                else "failed"
            )
            # Do not persist parser exceptions that might include document contents or credentials.
            error = (
                "optional parser dependency unavailable"
                if status == "needs_parser"
                else "parsing failed; original retained; retry after checking file/parser"
            )
            with self.store.transaction() as db:
                db.execute(
                    "UPDATE knowledge_documents SET status=?,error=?,lease=0,updated=? WHERE id=? AND status='processing' AND lease=?",
                    (status, error, self.engine.clock(), row["id"], lease),
                )
                if status == "blocked":
                    db.execute(
                        "UPDATE knowledge_documents SET name='',title='',description='',origin='{}',error='memory controls applied during parsing' WHERE id=? AND status='blocked'",
                        (row["id"],),
                    )
            if status == "blocked":
                with self.store.knowledge_dispatch():
                    self.purge_redacted()
        return True

    def get(self, p, id):
        with self.store.transaction() as db:
            row = self._get(db, p, id)
            return {
                **self._info(db, row),
                "locations": [
                    dict(r)
                    for r in db.execute(
                        "SELECT memory_id,ordinal,page,method FROM knowledge_chunks WHERE document_id=? ORDER BY ordinal",
                        (id,),
                    )
                ],
            }

    def list(self, p, *, query="", page=1, size=30, historical=False):
        if not 1 <= size <= 100 or page < 1:
            raise ValueError("invalid pagination")
        with self.store.transaction() as db:
            rows = db.execute(
                "SELECT * FROM knowledge_documents WHERE principal=? ORDER BY created DESC,id",
                (p.key,),
            ).fetchall()
            visible = [r for r in rows if self._allowed(db, p, r)]
            # Latest visible version per filename; history is explicitly requested.
            if not historical:
                latest = {}
                for row in visible:
                    if (
                        row["family"] not in latest
                        or row["version"] > latest[row["family"]]["version"]
                    ):
                        latest[row["family"]] = row
                visible = list(latest.values())
            if query:
                visible = [
                    r
                    for r in visible
                    if query.casefold()
                    in (r["name"] + " " + r["title"] + " " + r["description"]).casefold()
                ]
            start = (page - 1) * size
            return {
                "items": [self._info(db, r) for r in visible[start : start + size]],
                "total": len(visible),
                "page": page,
                "size": size,
            }

    def read(self, p, id, *, start=0, size=5):
        if start < 0 or not 1 <= size <= 20:
            raise ValueError("invalid chunk range")
        with self.store.transaction() as db:
            row = self._get(db, p, id)
            items = []
            for item in db.execute(
                "SELECT k.*,m.* FROM knowledge_chunks k JOIN memories m ON m.id=k.memory_id WHERE k.document_id=? AND k.ordinal>=? ORDER BY k.ordinal LIMIT ?",
                (id, start, size),
            ):
                if self.engine._fact_allowed(db, p, item):
                    items.append(
                        {
                            "memory_id": item["memory_id"],
                            "chunk": item["ordinal"],
                            "page": item["page"],
                            "method": item["method"],
                            "text": item["content"],
                        }
                    )
            return {
                "document": self._info(db, row),
                "items": items,
                "next": start + len(items)
                if start + len(items) < self._info(db, row)["chunks"]
                else None,
            }

    def original(self, p, id):
        # Caller gets a bounded byte copy while deletion and access are serialized.
        with self.store.knowledge_dispatch(), self.store.transaction() as db:
            row = self._get(db, p, id)
            try:
                data = self.blob(id).read_bytes()
            except FileNotFoundError as exc:
                raise KeyError("original file is unavailable") from exc
            return self._info(db, row), data

    def search(self, p, query, *, limit_items=10, historical=False):
        if not query.strip() or len(query) > 2000 or not 1 <= limit_items <= 50:
            raise ValueError("invalid search")
        q = terms(query)
        vector = None
        if self.engine.model and getattr(self.engine.model, "available", False):
            try:
                vector = self.engine.model.embed(limit(query, 256))
            except RuntimeError:
                pass
        with self.store.transaction() as db:
            candidates = {}
            if q:
                match = " OR ".join('"' + t.replace('"', '""') + '"' for t in sorted(q)[:64])
                for r in db.execute(
                    "SELECT m.* FROM memory_fts f JOIN memories m ON m.id=f.id JOIN knowledge_chunks k ON k.memory_id=m.id WHERE memory_fts MATCH ? AND m.principal=? ORDER BY bm25(memory_fts) LIMIT 100",
                    (match, p.key),
                ):
                    candidates[r["id"]] = decode(dict(r))
            for r in db.execute(
                "SELECT m.* FROM memories m JOIN knowledge_chunks k ON k.memory_id=m.id JOIN knowledge_documents d ON d.id=k.document_id WHERE m.principal=? AND (instr(lower(d.name||' '||d.title),lower(?))>0 OR (m.embedding IS NOT NULL AND m.embedding_space=?)) ORDER BY m.updated DESC LIMIT 2000",
                (p.key, query, getattr(self.engine.model, "space", None) if vector else None),
            ):
                candidates[r["id"]] = decode(dict(r))
            hits = []
            docs = {}
            latest = {
                r["family"]: r["v"]
                for r in db.execute(
                    "SELECT family,max(version) v FROM knowledge_documents WHERE principal=? AND status IN ('ready','partial') GROUP BY family",
                    (p.key,),
                )
            }
            for r in candidates.values():
                if r["status"] not in (
                    {"active", "superseded"} if historical else {"active"}
                ) or not self.engine._fact_allowed(db, p, r):
                    continue
                link = db.execute(
                    "SELECT * FROM knowledge_chunks WHERE memory_id=?", (r["id"],)
                ).fetchone()
                if link["document_id"] not in docs:
                    try:
                        docs[link["document_id"]] = self._get(db, p, link["document_id"])
                    except KeyError:
                        docs[link["document_id"]] = None
                doc = docs[link["document_id"]]
                if doc is None or not historical and doc["version"] != latest.get(doc["family"]):
                    continue
                lexical = len(
                    q & terms(r["content"] + " " + doc["name"] + " " + doc["title"])
                ) / max(1, len(q))
                cosine = 0
                v = r["embedding"]
                if (
                    vector
                    and v
                    and len(vector) == len(v)
                    and r["embedding_space"] == getattr(self.engine.model, "space", None)
                ):
                    import math

                    norm = math.sqrt(sum(x * x for x in vector) * sum(x * x for x in v))
                    cosine = (
                        sum(a * b for a, b in zip(vector, v, strict=True)) / norm if norm else 0
                    )
                if not lexical and cosine < 0.35:
                    continue
                hits.append(
                    {
                        "document_id": doc["id"],
                        "name": doc["name"],
                        "title": doc["title"],
                        "version": doc["version"],
                        "memory_id": r["id"],
                        "chunk": link["ordinal"],
                        "page": link["page"],
                        "method": link["method"],
                        "text": r["content"],
                        "score": round(max(lexical, cosine), 4),
                    }
                )
            hits.sort(key=lambda r: r["score"], reverse=True)
            return {
                "items": hits[:limit_items],
                "truncated": len(hits) > limit_items,
                "data_is_untrusted": True,
            }

    def retry(self, p, id):
        with self.store.transaction() as db:
            row = self._get(db, p, id)
            if row["status"] not in {"failed", "needs_parser", "empty"}:
                raise ValueError("only failed or unindexed documents can be retried")
            if db.execute(
                "SELECT 1 FROM exclusions WHERE principal=? AND state!='active'", (p.key,)
            ).fetchone():
                raise PermissionError("import paused by exclusions")
            db.execute(
                "UPDATE knowledge_documents SET status='queued',error='',epoch=?,lease=0 WHERE id=?",
                (self.engine._epoch(db, p), id),
            )
            Store.changed(db)
        return {"id": id, "status": "queued"}

    def embed_one(self):
        model = self.engine.model
        if (
            not model
            or not getattr(model, "available", False)
            or not getattr(self.engine.config, "embedding_model", "")
        ):
            return False
        now = self.engine.clock()
        if now < getattr(self.engine, "_knowledge_embedding_retry", 0):
            return False
        with self.store.transaction() as db:
            rows = db.execute(
                "SELECT m.* FROM memories m JOIN knowledge_chunks k ON k.memory_id=m.id JOIN knowledge_documents d ON d.id=k.document_id WHERE m.status='active' AND d.status IN ('ready','partial') AND (m.embedding IS NULL OR m.embedding_space IS NOT ?) ORDER BY m.created",
                (model.space,),
            ).fetchall()
            row = None
            for candidate in rows:
                parts = candidate["principal"].split("\x1f")
                principal = Principal(*parts, parts[0] if parts[1] == "owner" else parts[2])
                if self.engine._fact_allowed(db, principal, candidate):
                    row = dict(candidate)
                    break
            if row is None:
                return False
        try:
            vector = model.embed(row["content"])
        except RuntimeError:
            self.engine._knowledge_embedding_retry = now + 60
            return False
        if vector:
            parts = row["principal"].split("\x1f")
            p = Principal(*parts, parts[0] if parts[1] == "owner" else parts[2])
            with self.store.transaction() as db:
                current = db.execute(
                    "SELECT * FROM memories WHERE id=? AND status='active'", (row["id"],)
                ).fetchone()
                if (
                    current
                    and current["content"] == row["content"]
                    and self.engine._fact_allowed(db, p, current)
                ):
                    db.execute(
                        "UPDATE memories SET embedding=?,embedding_space=? WHERE id=?",
                        (compact(vector), model.space, row["id"]),
                    )
                    Store.changed(db)
        return bool(vector)

    def purge_redacted(self):
        """Called while the profile's knowledge file lock is held."""
        for row in self.store.rows(
            "SELECT id FROM knowledge_documents WHERE status IN ('redacted','blocked')"
        ):
            self.blob(row["id"]).unlink(missing_ok=True)

    def reconcile(self):
        """Remove uncommitted files after a crash; keep indexed data if an original is missing."""
        with self.store.knowledge_dispatch(), self.store.transaction() as db:
            documents = {
                r["id"]: r["status"]
                for r in db.execute("SELECT id,status FROM knowledge_documents")
            }
            removed = 0
            for path in (self.root / "blobs").glob("*/*"):
                if path.is_file() and (
                    path.name.startswith(".upload-")
                    or re.fullmatch(r"[a-f0-9]{64}\.blob", path.name)
                    and (
                        path.stem not in documents
                        or documents[path.stem] in {"redacted", "blocked"}
                    )
                ):
                    path.unlink()
                    removed += 1
            missing = sum(
                not self.blob(id).is_file()
                for id, status in documents.items()
                if status not in {"redacted", "blocked"}
            )
            return {"uncommitted_files_removed": removed, "missing_originals": missing}

    def backup(self, destination):
        destination = Path(destination)
        if destination.exists():
            raise ValueError("backup destination already exists")
        if self.store.path.parent.resolve() in destination.resolve().parents:
            raise ValueError("backup must be outside the live GraMem directory")
        with self.store.knowledge_dispatch(), self.store.transaction() as db:
            destination.parent.mkdir(parents=True, exist_ok=True)
            fd, tmp = tempfile.mkstemp(dir=destination.parent, prefix=".knowledge-backup-")
            os.close(fd)
            try:
                # A second read connection sees the last committed state; writers are blocked.
                import sqlite3

                with tempfile.TemporaryDirectory() as folder:
                    sqlite = Path(folder) / "memory.sqlite"
                    # sqlite3 connection context managers commit but do not close.
                    source = sqlite3.connect(str(self.store.path))
                    target = sqlite3.connect(str(sqlite))
                    try:
                        source.backup(target)
                    finally:
                        target.close()
                        source.close()
                    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                        archive.write(sqlite, "memory.sqlite")
                        documents = [
                            dict(r)
                            for r in db.execute(
                                "SELECT id,sha256,size FROM knowledge_documents WHERE status NOT IN ('redacted','blocked')"
                            )
                        ]
                        for row in documents:
                            archive.write(
                                self.blob(row["id"]),
                                f"knowledge/blobs/{row['id'][:2]}/{row['id']}.blob",
                            )
                        from .chat_audio import ChatAudio

                        audio = [dict(r) for r in db.execute("SELECT id,sha256,size FROM chat_audio WHERE state='stored'")]
                        for row in audio:
                            data = ChatAudio(self.engine).blob(row['id']).read_bytes()
                            if hashlib.sha256(data).hexdigest() != row['sha256'] or len(data) != row['size']:
                                raise RuntimeError('audio original integrity failure')
                            archive.writestr(f"chats/audio/{row['id'][:2]}/{row['id']}.blob", data)
                        archive.writestr(
                            "manifest.json",
                            compact(
                                {"schema": db.execute('PRAGMA user_version').fetchone()[0],
                                 "documents": documents, "chat_audio": audio, "graph_backup_required": True}
                            ),
                        )
                # Do not overwrite a concurrent backup.
                os.link(tmp, destination)
            finally:
                Path(tmp).unlink(missing_ok=True)
        return {
            "backup": str(destination),
            "documents": len(documents),
            "chat_audio": len(audio),
            "graph_backup_required": True,
        }
