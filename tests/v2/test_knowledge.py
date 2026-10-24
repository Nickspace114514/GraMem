import asyncio
import io
import json
import sqlite3
import types
import zipfile
from pathlib import Path

import pytest
from test_engine import CONTACT, GROUP, OWNER
from test_provider import make_provider  # noqa: F401, F811

from gramem.config import Config
from gramem.engine import Engine
from gramem.knowledge import Knowledge


@pytest.fixture
def library(tmp_path):
    engine = Engine(tmp_path / "gramem" / "memory.sqlite")
    yield Knowledge(engine)
    engine.close()


def save(library, text="中文知识：关系图使用 FalkorDB。", name="guide.md", **kwargs):
    return library.ingest(OWNER, data=text.encode(), filename=name, **kwargs)


def test_original_roundtrip_scope_and_persistent_reopen(library):
    doc = save(library)
    assert library.original(OWNER, doc["id"])[1] == "中文知识：关系图使用 FalkorDB。".encode()
    assert library.blob(doc["id"]).stat().st_mode & 0o777 == 0o600
    with pytest.raises(KeyError):
        library.original(CONTACT, doc["id"])
    assert not library.list(GROUP)["items"]
    reopened = Engine(library.store.path)
    try:
        hit = Knowledge(reopened).search(OWNER, "关系图")["items"][0]
        assert hit["document_id"] == doc["id"] and hit["page"] == 1
    finally:
        reopened.close()


def test_identical_bytes_deduplicate_but_other_identity_is_independent(library):
    a = save(library)
    b = save(library, name="renamed.md")
    assert b["id"] == a["id"] and b["duplicate"]
    contact = library.ingest(
        CONTACT, data="中文知识：关系图使用 FalkorDB。".encode(), filename="guide.md"
    )
    assert contact["id"] != a["id"]
    library.engine.exclude(OWNER, "document:" + a["id"])
    assert library.original(CONTACT, contact["id"])[1]


def test_versions_default_recall_uses_new_and_history_keeps_original(library):
    first = save(library, "旧方案使用苹果架构")
    second = save(library, "新方案使用橙子架构")
    assert second["version"] == 2 and second["previous"] == first["id"]
    assert library.list(OWNER)["total"] == 1
    assert library.list(OWNER, historical=True)["total"] == 2
    assert not library.search(OWNER, "苹果")["items"]
    assert library.search(OWNER, "苹果", historical=True)["items"]
    assert "旧方案" not in library.engine.recall(OWNER, "架构")["context"]
    assert "新方案" in library.engine.recall(OWNER, "架构")["context"]
    assert library.original(OWNER, first["id"])[1] == "旧方案使用苹果架构".encode()


def test_bad_new_version_does_not_remove_previous_index(library):
    first = save(library, "原方案使用苹果架构")
    bad = library.ingest(OWNER, data=b"not a pdf", filename="guide.pdf")
    assert bad["status"] == "failed" and library.original(OWNER, bad["id"])[1]
    assert library.search(OWNER, "苹果")["items"][0]["document_id"] == first["id"]


def test_document_suppression_hides_original_chunks_and_derived_facts(library):
    doc = save(library)
    memory = doc["ids"][0]
    derived = library.engine.remember(
        OWNER, "总结：使用 FalkorDB", action="MERGE", targets=[memory], slot="summary"
    )
    library.engine.exclude(OWNER, "document:" + doc["id"], state="suppressed")
    assert not library.search(OWNER, "FalkorDB")["items"]
    assert not library.engine.recall(OWNER, "FalkorDB")["context"]
    assert derived not in [r["id"] for r in library.engine.list(OWNER)]
    with pytest.raises(KeyError):
        library.original(OWNER, doc["id"])
    library.engine.reactivate(OWNER, "document:" + doc["id"])
    assert library.original(OWNER, doc["id"])[1]
    assert "FalkorDB" in library.engine.recall(OWNER, "FalkorDB")["context"]


def test_forgetting_any_chunk_erases_file_and_all_chunks(library):
    doc = save(library, "长期知识库使用独立存储。" * 200)
    row = library.engine.get(OWNER, doc["ids"][1])
    library.engine.exclude(OWNER, row["slot"])
    assert not library.blob(doc["id"]).exists()
    assert not library.list(OWNER)["items"]
    assert all(r["content"] == "" for r in library.store.rows("SELECT content FROM memories"))
    with pytest.raises(PermissionError):
        save(library, "长期知识库使用独立存储。" * 200)


def test_wildcard_forget_erases_queued_original_without_chunks(library):
    doc = save(library, defer=True)
    assert doc["status"] == "queued" and not doc["ids"]
    library.engine.exclude(OWNER, "*")
    assert not library.blob(doc["id"]).exists()
    assert not library.process_one()


def test_durable_queue_claim_recovers_after_crash(library):
    doc = save(library, defer=True)
    with library.store.transaction() as db:
        db.execute(
            "UPDATE knowledge_documents SET status='processing',lease=1 WHERE id=?", (doc["id"],)
        )
    assert library.process_one()
    assert library.get(OWNER, doc["id"])["status"] == "ready"
    assert not library.process_one()


def test_reconcile_cleans_crash_orphans_and_preserves_committed_files(library):
    doc = save(library)
    orphan = library.blob("a" * 64)
    orphan.parent.mkdir(parents=True, exist_ok=True)
    orphan.write_bytes(b"uncommitted")
    temporary = orphan.parent / ".upload-aborted"
    temporary.write_bytes(b"partial")
    assert library.reconcile() == {"uncommitted_files_removed": 2, "missing_originals": 0}
    assert library.original(OWNER, doc["id"])[1]
    library.blob(doc["id"]).unlink()
    assert library.reconcile()["missing_originals"] == 1
    assert library.search(OWNER, "关系图")["items"]
    with pytest.raises(KeyError):
        library.original(OWNER, doc["id"])


def test_race_delete_during_parse_does_not_recreate_content(library, monkeypatch):
    from gramem import knowledge

    doc = save(library, defer=True)
    original = knowledge.extract

    def erase(data, suffix, **kwargs):
        library.engine.exclude(OWNER, "document:" + doc["id"])
        return original(data, suffix)

    monkeypatch.setattr(knowledge, "extract", erase)
    assert library.process_one()
    assert not library.blob(doc["id"]).exists()
    assert not library.store.rows("SELECT * FROM sources")
    assert not library.store.rows("SELECT * FROM memories")


def test_ocr_unavailable_retains_image_and_caption_is_searchable(library, monkeypatch):
    from PIL import Image

    from gramem import knowledge

    def missing(path, languages=None):
        raise ImportError("missing")

    monkeypatch.setattr(knowledge, "_ocr", missing)
    buffer = io.BytesIO()
    Image.new("RGB", (40, 40), "white").save(buffer, format="PNG")
    doc = library.ingest(
        OWNER, data=buffer.getvalue(), filename="sample.png", description="测试参考图"
    )
    assert doc["status"] == "partial"
    hit = library.search(OWNER, "参考图")["items"][0]
    assert hit["method"] == "description"
    assert library.original(OWNER, doc["id"])[1] == buffer.getvalue()


def test_scanned_pdf_page_location(library, monkeypatch):
    from pypdf import PdfWriter

    from gramem import knowledge

    pdf = PdfWriter()
    pdf.add_blank_page(width=100, height=100)
    pdf.add_blank_page(width=100, height=100)
    buffer = io.BytesIO()
    pdf.write(buffer)
    monkeypatch.setattr(knowledge.shutil, "which", lambda name: "/usr/bin/" + name)

    def render(args, timeout=60):
        return ""

    monkeypatch.setattr(knowledge, "_run", render)
    monkeypatch.setattr(knowledge, "_ocr", lambda path, languages=None: "扫描页的中文知识")
    doc = library.ingest(OWNER, data=buffer.getvalue(), filename="scanned.pdf")
    assert doc["pages"] == 2 and doc["status"] == "ready"
    assert [r["page"] for r in library.read(OWNER, doc["id"])["items"]] == [1, 2]
    assert all(r["method"] == "ocr" for r in library.read(OWNER, doc["id"])["items"])


def test_docx_paragraph_and_table_without_invented_pages(library):
    from docx import Document

    doc = Document()
    doc.add_paragraph("文档知识库段落")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "表格知识"
    buffer = io.BytesIO()
    doc.save(buffer)
    result = library.ingest(OWNER, data=buffer.getvalue(), filename="guide.docx")
    text = library.read(OWNER, result["id"])["items"][0]
    assert "表格知识" in text["text"] and text["page"] == 0


def test_backup_contains_sqlite_and_every_original(library, tmp_path):
    doc = save(library)
    target = tmp_path / "backup.zip"
    assert library.backup(target)["documents"] == 1
    with zipfile.ZipFile(target) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        assert manifest["schema"] == 8
        assert (
            archive.read(f"knowledge/blobs/{doc['id'][:2]}/{doc['id']}.blob")
            == library.original(OWNER, doc["id"])[1]
        )
        database = tmp_path / "restored.sqlite"
        database.write_bytes(archive.read("memory.sqlite"))
        with sqlite3.connect(database) as db:
            assert db.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
            assert db.execute("SELECT count(*) FROM knowledge_documents").fetchone()[0] == 1
    assert target.stat().st_mode & 0o777 == 0o600
    with pytest.raises(ValueError):
        library.backup(target)


def test_secret_and_bad_files_rejected_without_original(library):
    with pytest.raises(ValueError):
        save(library, "api_key=do-not-store-me")
    with pytest.raises(ValueError):
        save(library, name="program.exe")
    with pytest.raises(PermissionError):
        library.ingest(GROUP, data=b"private", filename="guide.txt")
    assert not library.store.rows("SELECT * FROM knowledge_documents")


def test_retry_invalid_pdf_keeps_original(library):
    doc = library.ingest(OWNER, data=b"not a pdf", filename="broken.pdf")
    assert doc["status"] == "failed"
    assert library.retry(OWNER, doc["id"])["status"] == "queued"
    assert library.original(OWNER, doc["id"])[1] == b"not a pdf"


def test_embedding_exclusions_checked_before_external_call(library):
    doc = save(library)

    class Model:
        available = True
        space = "test-space"
        calls = 0

        def embed(self, text):
            self.calls += 1
            return [1.0, 0.0]

        def close(self):
            pass

    model = Model()
    library.engine.model = model
    library.engine.config.embedding_model = "test-model"
    library.engine.exclude(OWNER, "document:" + doc["id"], state="suppressed")
    assert not library.embed_one() and model.calls == 0
    library.engine.reactivate(OWNER, "document:" + doc["id"])
    assert library.embed_one() and model.calls == 1
    assert library.search(OWNER, "other wording")["items"][0]["document_id"] == doc["id"]


def test_out_of_order_workers_keep_newest_version_active(library):
    first = save(library, "旧版本苹果方案", defer=True)
    second = save(library, "新版本橙子方案", defer=True)
    library.process_one(id=second["id"])
    library.process_one(id=first["id"])
    assert not library.search(OWNER, "苹果")["items"]
    assert library.search(OWNER, "苹果", historical=True)["items"]
    assert "旧版本" not in library.engine.recall(OWNER, "方案")["context"]


def test_private_tools_and_export_shared_scope_denied(make_provider):  # noqa: F811
    provider = make_provider(chat_type="direct")
    principal = provider._current
    kb = Knowledge(provider.engine)
    doc = kb.ingest(principal, data="用户的知识资料".encode(), filename="notes.md")
    provider.on_turn_start(1, "查找知识资料", author_id="verified")
    result = json.loads(
        provider.handle_tool_call("gramem_knowledge", {"action": "search", "query": "知识资料"})
    )
    assert result["items"]
    exported = json.loads(
        provider.handle_tool_call("gramem_knowledge", {"action": "export", "id": doc["id"]})
    )
    assert Path(exported["path"]).read_bytes() == "用户的知识资料".encode()
    provider.chat_type = "channel"
    provider.on_turn_start(2, "读取我的文件", author_id="verified")
    assert not json.loads(
        provider.handle_tool_call("gramem_knowledge", {"action": "search", "query": "知识资料"})
    )["items"]
    assert "error" in json.loads(
        provider.handle_tool_call("gramem_knowledge", {"action": "export", "id": doc["id"]})
    )


def test_gateway_real_metadata_only_authorized_owner(tmp_path, monkeypatch):
    from gramem.knowledge_gateway import archive_inbound

    home = tmp_path / "home"
    (home / "gramem").mkdir(parents=True)
    (home / "gramem/config.json").write_text(
        json.dumps(
            {
                "owner_id": "owner",
                "aliases": {"teams:verified": "owner"},
                "capture_owner_in_groups": True,
            }
        )
    )
    constants = types.ModuleType("hermes_constants")
    constants.get_hermes_home = lambda: home
    monkeypatch.setitem(__import__("sys").modules, "hermes_constants", constants)
    source = types.SimpleNamespace(
        user_id="verified",
        platform=types.SimpleNamespace(value="teams"),
        chat_type="dm",
        chat_id="chat",
        is_bot=False,
    )
    cache = home / "cache/documents"
    cache.mkdir(parents=True)
    attachment = cache / "doc_123456abcdef_notes.txt"
    attachment.write_text("真实附件资料")
    event = types.SimpleNamespace(
        source=source, media_urls=[str(attachment)], text="我的资料", message_id="m1"
    )
    gateway = types.SimpleNamespace(
        _is_user_authorized_for_source=lambda s: False,
        _resolve_profile_home_for_source=lambda s: home,
    )
    asyncio.run(archive_inbound(event, gateway))
    assert not (home / "gramem/memory.sqlite").exists()
    gateway._is_user_authorized_for_source = lambda s: True
    source.user_id = "guest"
    asyncio.run(archive_inbound(event, gateway))
    assert not (home / "gramem/memory.sqlite").exists()
    source.user_id = "verified"
    asyncio.run(archive_inbound(event, gateway))
    e = Engine(home / "gramem/memory.sqlite", Config(owner_id="owner"))
    try:
        docs = Knowledge(e).list(OWNER)["items"]
        assert len(docs) == 1 and docs[0]["name"] == "notes.txt"
        assert docs[0]["status"] == "queued"
        assert docs[0]["origin"]["message_id"] == "m1"
        assert Knowledge(e).original(OWNER, docs[0]["id"])[1] == "真实附件资料".encode()
    finally:
        e.close()
